#!/usr/bin/env python3
"""
UAT Script Generator

Reads an integration workflow export (Tray-style JSON) and, optionally, the signed
requirements spec for that integration. Produces a customer-ready UAT workbook:

  Overview · Pre-UAT Validation · Field Mappings · Test Cases · Test Data ·
  Issue Log · Sign-off · AI Instructions

Usage:
  python generate_uat.py --workflow wf.json [--spec spec.json] [--out uat.xlsx]

Prints a JSON summary to stdout (validation results + test-case counts) so the
calling agent can brief the user before they open the workbook.
"""
import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# --------------------------------------------------------------------------------------
# 1. Parse the workflow
# --------------------------------------------------------------------------------------
UTILITY_CONNECTORS = {"loop", "boolean-condition", "data-mapper", "error-handler", "script", "branch"}
HRIS_CONNECTORS = {"workday", "successfactors", "successfactors-ec", "oracle-hcm", "bamboohr", "adp", "ukg"}
LMS_CONNECTORS = {"successfactors-lms", "cornerstone", "oracle-learning", "workday-learning", "docebo"}


def schedule_minutes(value):
    """'15m' -> 15, '2h' -> 120, 'daily' -> 1440, 'weekly' -> 10080."""
    if value is None:
        return None
    v = str(value).strip().lower()
    named = {"hourly": 60, "daily": 1440, "weekly": 10080}
    if v in named:
        return named[v]
    m = re.fullmatch(r"(\d+)\s*(m|min|h|hr|d)", v)
    if not m:
        return None
    n, unit = int(m.group(1)), m.group(2)
    return n * {"m": 1, "min": 1, "h": 60, "hr": 60, "d": 1440}[unit]


def human_schedule(value):
    if value is None:
        return "not set"
    mins = schedule_minutes(value)
    if mins is None:
        return str(value)
    if mins % 1440 == 0:
        return "daily" if mins == 1440 else f"every {mins // 1440} days"
    if mins % 60 == 0:
        return "hourly" if mins == 60 else f"every {mins // 60} hours"
    return f"every {mins} minutes"


def parse_workflow(raw):
    wf = raw.get("workflow", raw)
    steps = wf.get("steps", [])
    business = [s for s in steps if s.get("connector") not in UTILITY_CONNECTORS]
    source = business[0] if business else {}
    target = business[-1] if len(business) > 1 else {}

    mappings, filters, retry = [], [], []
    for s in steps:
        props = s.get("properties", {})
        if s.get("connector") == "data-mapper":
            mappings.extend(props.get("mappings", []))
        if s.get("connector") == "boolean-condition":
            filters.extend(props.get("conditions", []))
        if s.get("connector") == "error-handler":
            retry = props.get("retry_on_status", [])

    tags = wf.get("tags", [])
    itype = next((t for t in tags if t in ("hris_user_sync", "lms_completions")), None)
    if not itype:  # infer from connectors
        if source.get("connector") in HRIS_CONNECTORS:
            itype = "hris_user_sync"
        elif target.get("connector") in LMS_CONNECTORS:
            itype = "lms_completions"
        else:
            itype = "unknown"
    env = next((t.split(":", 1)[1] for t in tags if t.startswith("env:")), "unspecified")

    return {
        "id": wf.get("id", ""),
        "title": wf.get("title", "Untitled workflow"),
        "customer": wf.get("customer", "Customer"),
        "type": itype,
        "env": env,
        "schedule": (wf.get("trigger", {}).get("schedule", {}) or {}).get("interval"),
        "source": source,
        "target": target,
        "mappings": mappings,
        "filters": filters,
        "retry": retry,
    }


def mapping_by_target(parsed):
    return {m["target"]: m for m in parsed["mappings"]}


# --------------------------------------------------------------------------------------
# 2. Validate the workflow against the signed requirements spec
# --------------------------------------------------------------------------------------
def validate(parsed, spec):
    checks = []

    def add(category, requirement, as_built, result, severity, action, field=None):
        checks.append({
            "id": f"V-{len(checks) + 1:02d}", "category": category, "requirement": requirement,
            "as_built": as_built, "result": result, "severity": severity, "action": action, "field": field,
        })

    if spec is None:
        return checks

    by_target = mapping_by_target(parsed)

    # Integration type
    st = spec.get("integration_type")
    add("Scope", f"Integration type: {st}", f"Workflow tagged: {parsed['type']}",
        "PASS" if st == parsed["type"] else "FAIL", "—" if st == parsed["type"] else "High",
        "—" if st == parsed["type"] else "Confirm the right workflow was exported.")

    # Unique key
    uk = spec.get("unique_key", {})
    if uk:
        match_on = parsed["target"].get("properties", {}).get("match_on") or uk.get("target")
        src = by_target.get(uk.get("target"), {}).get("source")
        ok = match_on == uk.get("target") and src == uk.get("source")
        add("Identity", f"Records match on {uk.get('source')} → {uk.get('target')}",
            f"Match on {match_on}; {uk.get('target')} ← {src or 'unmapped'}",
            "PASS" if ok else "FAIL", "—" if ok else "High",
            "—" if ok else "Fix the match key before any data loads; a wrong key creates duplicate users.",
            uk.get("target"))

    # Required mappings
    for req in spec.get("required_mappings", []):
        tgt, src = req["target"], req["source"]
        why = f" ({req['note']})" if req.get("note") else ""
        built = by_target.get(tgt)
        if not built:
            add("Field mapping", f"{tgt} ← {src}{why}", "Not mapped", "FAIL", "High",
                f"Add mapping {src} → {tgt} in the data mapper.", tgt)
        elif built["source"] != src:
            add("Field mapping", f"{tgt} ← {src}{why}", f"{tgt} ← {built['source']}", "FAIL", "High",
                f"Change source from {built['source']} to {src}.", tgt)
        else:
            add("Field mapping", f"{tgt} ← {src}", f"{tgt} ← {built['source']}", "PASS", "—", "—", tgt)

    # Mappings that were built but never agreed (only meaningful if the spec lists mappings at all)
    agreed = {r["target"] for r in spec.get("required_mappings", [])}
    if not agreed:
        add("Spec coverage", "Signed spec lists the field mappings", "Spec has no required_mappings",
            "WARN", "Medium", "Field mappings were not validated. Get the agreed mappings into the spec before UAT.")
    for m in parsed["mappings"] if agreed else []:
        if m["target"] not in agreed:
            add("Field mapping", "Not in signed spec", f"{m['target']} ← {m['source']}", "INFO", "Low",
                "Confirm with the customer that this field should sync, or remove it.", m["target"])

    # Filters
    built_filters = {f["field"].lower(): f for f in parsed["filters"]}
    for f in spec.get("filters", []):
        b = built_filters.get(f["field"].lower())
        want = f"{f['field']} {f['operator']} {', '.join(f['value'])}"
        if not b:
            add("Population filter", want, "No filter on this field", "FAIL", "High",
                "Add the filter; out-of-scope workers will otherwise be created.", f["field"])
            continue
        missing = sorted(set(f["value"]) - set(b.get("value", [])))
        extra = sorted(set(b.get("value", [])) - set(f["value"]))
        got = f"{b['field']} {b['operator']} {', '.join(b.get('value', []))}"
        if b["operator"] != f["operator"] or missing:
            detail = f"Missing value(s): {', '.join(missing)}. " if missing else ""
            add("Population filter", want, got, "FAIL", "High",
                f"{detail}Update the filter to match the signed spec.", f["field"])
        elif extra:
            add("Population filter", want, got, "WARN", "Medium",
                f"Workflow also filters {', '.join(extra)}; confirm with the customer.", f["field"])
        else:
            add("Population filter", want, got, "PASS", "—", "—", f["field"])

    # Schedule
    if spec.get("schedule"):
        want, got = schedule_minutes(spec["schedule"]), schedule_minutes(parsed["schedule"])
        label_w, label_g = human_schedule(spec["schedule"]), human_schedule(parsed["schedule"])
        if parsed["schedule"] is None:
            label_g = "on no schedule (trigger not set)"
        if want == got:
            add("Schedule", f"Runs {label_w}", f"Runs {label_g}", "PASS", "—", "—")
        elif got is not None and want is not None and got < want:
            add("Schedule", f"Runs {label_w}", f"Runs {label_g}", "WARN", "Medium",
                f"Faster cadence is fine for UAT. Revert to {label_w} before production cutover.")
        else:
            add("Schedule", f"Runs {label_w}", f"Runs {label_g}", "FAIL", "Medium",
                f"Set the trigger to {label_w}.")

    # Lifecycle / status handling
    lc = spec.get("lifecycle")
    if lc:
        status_map = by_target.get("status", {}).get("value_map", {})
        for src_val, want in lc.items():
            if src_val == "missing_from_feed":
                got = parsed["target"].get("properties", {}).get("on_missing_from_feed", "not set")
                label = "Worker missing from feed"
            else:
                got = status_map.get(src_val, "not mapped")
                label = f"Status '{src_val}'"
            ok = got == want
            add("Lifecycle", f"{label} → {want}", f"{label} → {got}", "PASS" if ok else "FAIL",
                "—" if ok else "High", "—" if ok else f"Set '{src_val}' handling to {want}.", "status")

    # Timezone on datetime transforms
    tz = spec.get("timezone")
    if tz:
        for m in parsed["mappings"]:
            if str(m.get("transform", "")).startswith("datetime"):
                got = m.get("timezone")
                if got == tz:
                    add("Dates & time zones", f"{m['target']} converted in {tz}", f"Converted in {got}",
                        "PASS", "—", "—", m["target"])
                else:
                    add("Dates & time zones", f"{m['target']} converted in {tz}",
                        f"Converted in {got or 'unspecified (platform default)'}", "FAIL", "High",
                        f"Set timezone to {tz}. Late-evening completions will otherwise land on the "
                        f"next calendar day in the customer's system.", m["target"])

    # Error handling
    eh = spec.get("error_handling", {})
    if eh.get("retry_on_status"):
        missing = sorted(set(eh["retry_on_status"]) - set(parsed["retry"]))
        got = ", ".join(map(str, parsed["retry"])) or "none"
        want = ", ".join(map(str, eh["retry_on_status"]))
        if missing:
            note = " 429 = API rate limit; large first loads will drop records without it." if 429 in missing else ""
            add("Error handling", f"Retry on HTTP {want}", f"Retries on {got}", "FAIL", "High",
                f"Add retry on {', '.join(map(str, missing))}.{note}")
        else:
            add("Error handling", f"Retry on HTTP {want}", f"Retries on {got}", "PASS", "—", "—")

    # IDs follow spec order, so they stay the same between runs. Display order puts problems first.
    rank = {"FAIL": 0, "WARN": 1, "INFO": 2, "PASS": 3}
    checks.sort(key=lambda c: rank[c["result"]])  # stable sort: spec order kept within each group
    return checks


# --------------------------------------------------------------------------------------
# 3. Test data personas (all fictional)
# --------------------------------------------------------------------------------------
def domain_for(customer):
    slug = re.sub(r"[^a-z0-9]", "", customer.lower()) or "customer"
    return f"{slug}.example"


def hris_personas(customer):
    d = domain_for(customer)
    p = [
        ("P-01", "New hire", "HF-20417", "Renata", "Alcott", f"renata.alcott@{d}", "FULL_TIME", "Active", "2026-10-12", "2026-10-12", "Operations", "Dispatch Coordinator", "Joliet, IL", f"tomas.berg@{d}"),
        ("P-02", "Existing active worker", "HF-11802", "Devon", "Akintola", f"devon.akintola@{d}", "FULL_TIME", "Active", "2019-03-04", "2019-03-04", "Fleet Maintenance", "Diesel Technician II", "Gary, IN", f"lena.kovacs@{d}"),
        ("P-03", "Termination", "HF-14455", "Graham", "Pell", f"graham.pell@{d}", "FULL_TIME", "Terminated", "2021-06-14", "2021-06-14", "Sales", "Account Executive", "Chicago, IL", f"maya.shah@{d}"),
        ("P-04", "Rehire", "HF-09310", "Lucia", "Ferreira", f"lucia.ferreira@{d}", "FULL_TIME", "Active", "2026-10-05", "2017-08-21", "Warehouse", "Shift Lead", "Joliet, IL", f"tomas.berg@{d}"),
        ("P-05", "Leave of absence", "HF-12871", "Samir", "Haddad", f"samir.haddad@{d}", "FULL_TIME", "On Leave", "2020-01-13", "2020-01-13", "Finance", "AP Specialist", "Chicago, IL", f"maya.shah@{d}"),
        ("P-06", "Email / name change", "HF-15560", "Jordan", "Whitaker-Lee", f"jordan.whitaker-lee@{d}", "FULL_TIME", "Active", "2022-02-07", "2022-02-07", "HR", "HR Business Partner", "Chicago, IL", f"maya.shah@{d}"),
        ("P-07", "Excluded class (contractor)", "HF-C0381", "Ines", "Moreau", f"ines.moreau@{d}", "CONTRACTOR", "Active", "2026-09-01", "2026-09-01", "IT", "Network Consultant", "Remote", f"lena.kovacs@{d}"),
        ("P-08", "Excluded class (seasonal)", "HF-S1142", "Kwame", "Boateng", f"kwame.boateng@{d}", "SEASONAL", "Active", "2026-10-01", "2026-10-01", "Warehouse", "Seasonal Loader", "Gary, IN", f"tomas.berg@{d}"),
        ("P-09", "Duplicate email", "HF-16023", "Chris", "Nguyen", f"chris.nguyen@{d}", "FULL_TIME", "Active", "2023-05-15", "2023-05-15", "Operations", "Route Planner", "Joliet, IL", f"tomas.berg@{d}"),
        ("P-10", "Duplicate email (2nd record)", "HF-16024", "Christine", "Nguyen", f"chris.nguyen@{d}", "PART_TIME", "Active", "2026-09-22", "2026-09-22", "Customer Service", "CS Representative", "Remote", f"maya.shah@{d}"),
        ("P-11", "Special characters", "HF-17788", "Zoë", "Ó Briain-Søndergaard", f"zoe.obriain@{d}", "FULL_TIME", "Active", "2024-11-04", "2024-11-04", "Legal", "Contracts Manager", "Chicago, IL", f"maya.shah@{d}"),
        ("P-12", "Missing required field", "HF-18120", "Pat", "Delgado", "", "FULL_TIME", "Active", "2026-10-06", "2026-10-06", "Warehouse", "Forklift Operator", "Gary, IN", f"tomas.berg@{d}"),
        ("P-13", "Manager change", "HF-13002", "Aisha", "Rahman", f"aisha.rahman@{d}", "FULL_TIME", "Active", "2018-09-10", "2018-09-10", "Operations", "Terminal Manager", "Joliet, IL", f"lena.kovacs@{d}"),
    ]
    cols = ["Persona", "Scenario", "Employee_ID", "Preferred_First_Name", "Last_Name", "Work_Email",
            "Employee_Class", "Worker_Status", "Hire_Date", "Original_Hire_Date", "Department",
            "Job_Title", "Work_Location", "Manager_Email"]
    return cols, [dict(zip(cols, r)) for r in p]


def lms_personas(customer):
    p = [
        ("P-01", "Standard completion", "HF-11802", "Devon Akintola", "SAFE-101", "Forklift Safety Refresher", "OLT", "2026-10-08 14:22", "92"),
        ("P-02", "Late-evening completion", "HF-12871", "Samir Haddad", "COMP-220", "Anti-Bribery Essentials", "OLT", "2026-10-08 23:41", "88"),
        ("P-03", "Non-OLT course", "HF-15560", "Jordan Whitaker-Lee", "ILT-040", "Manager Workshop (in person)", "ILT", "2026-10-07 16:05", ""),
        ("P-04", "No matching LMS student", "HF-20417", "Renata Alcott", "SAFE-101", "Forklift Safety Refresher", "OLT", "2026-10-08 09:10", "100"),
        ("P-05", "Duplicate completion", "HF-13002", "Aisha Rahman", "SAFE-101", "Forklift Safety Refresher", "OLT", "2026-10-08 10:30", "95"),
        ("P-06", "Retake on a later date", "HF-13002", "Aisha Rahman", "SAFE-101", "Forklift Safety Refresher", "OLT", "2026-10-09 08:15", "98"),
        ("P-07", "Course missing from LMS catalog", "HF-17788", "Zoë Ó Briain-Søndergaard", "NEW-999", "Pilot Course (not in catalog)", "OLT", "2026-10-08 11:00", "75"),
    ]
    cols = ["Persona", "Scenario", "learner.external_id", "Learner name", "course.external_course_id",
            "Course title", "course.course_type", "completed_at", "status", "score"]
    return cols, [dict(zip(cols, r[:8] + ("completed",) + r[8:])) for r in p]


# --------------------------------------------------------------------------------------
# 4. Test cases — generated from what the workflow actually does
# --------------------------------------------------------------------------------------
def find_check(checks, field=None, category=None):
    """All open (FAIL/WARN) checks matching a field and/or category. A test case can be blocked by several."""
    return [c for c in checks
            if (field is None or c["field"] == field) and (category is None or c["category"] == category)
            and c["result"] in ("FAIL", "WARN")]


def hris_cases(parsed, spec, checks):
    by_t = mapping_by_target(parsed)
    spec_targets = {r["target"] for r in (spec or {}).get("required_mappings", [])}
    has = lambda t: t in by_t or t in spec_targets  # test what was agreed, even if not yet built
    target_name = (spec or {}).get("target_system", "the target platform")
    src_name = (spec or {}).get("source_system", "the source system")
    cases = []

    def case(title, cat, persona, pre, steps, expected, checks_=None):
        cases.append({"title": title, "category": cat, "persona": persona, "pre": pre,
                      "steps": steps, "expected": expected, "checks": checks_ or []})

    run = f"Run the sync (or wait for the next scheduled run)."
    case("New hire is created", "Core", "P-01",
         f"P-01 exists in {src_name} with a future-dated or current hire date; not yet in {target_name}.",
         f"1. Confirm P-01 is in the source report.\n2. {run}\n3. Search {target_name} for external_id HF-20417.",
         f"User created with external_id, email, name and every mapped field populated (check against Field Mappings tab).")
    case("All mapped fields populate correctly", "Core", "P-02",
         "P-02 already active in both systems.",
         f"1. Export P-02's record from {src_name}.\n2. Open the user in {target_name}.\n3. Compare each field listed on the Field Mappings tab.",
         "Every field matches the source value after transforms; mark each row Verified on the Field Mappings tab.",
         find_check(checks, category="Field mapping"))
    case("Attribute update propagates", "Core", "P-02",
         "P-02 active and synced.",
         f"1. In {src_name}, change P-02's Job_Title to 'Diesel Technician III'.\n2. {run}\n3. Check the user in {target_name}.",
         "Job title updates on the existing user. No duplicate user is created.")
    if "status" in by_t or (spec or {}).get("lifecycle"):
        case("Termination deactivates the user", "Edge", "P-03",
             "P-03 active in both systems.",
             f"1. Terminate P-03 in {src_name} (effective today).\n2. {run}\n3. Check P-03 in {target_name}.",
             "User is set to inactive and stops receiving content. Completion history is retained.",
             find_check(checks, field="status"))
        case("Leave of absence keeps the user active", "Edge", "P-05",
             "P-05 active in both systems.",
             f"1. Place P-05 on leave in {src_name}.\n2. {run}\n3. Check P-05 in {target_name}.",
             "User remains active (per signed spec). Status field reflects leave if exposed.",
             find_check(checks, field="status"))
    if has("hire_date"):
        case("Rehire reactivates the original user", "Edge", "P-04",
             "P-04 was terminated in 2024 and has an inactive user in the target. Rehired 2026-10-05 under the same Employee_ID.",
             f"1. Confirm P-04 shows Hire_Date 2026-10-05 and Original_Hire_Date 2017-08-21 in {src_name}.\n2. {run}\n3. Check P-04 in {target_name}.",
             "Existing user is reactivated (no duplicate) and hire_date = 2026-10-05 so new-hire content is assigned.",
             find_check(checks, field="hire_date"))
    case("Email change updates, does not duplicate", "Edge", "P-06",
         "P-06 active and synced with jordan.whitaker-lee@….",
         f"1. Change P-06's Work_Email to jordan.lee@{domain_for(parsed['customer'])} in {src_name}.\n2. {run}\n3. Search {target_name} by external_id HF-15560.",
         "Same user now shows the new email. No second account. SSO login still works.",
         find_check(checks, field="external_id"))
    case("Name change propagates", "Edge", "P-06",
         "P-06 active and synced.",
         f"1. Change P-06's Last_Name to 'Lee' in {src_name}.\n2. {run}",
         "Last name updates on the existing user; display name refreshes in reports.")
    built_vals = {v for bf in parsed["filters"] for v in bf.get("value", [])}
    for flt in (spec or {}).get("filters", []) or parsed["filters"]:
        for val, pid in (("CONTRACTOR", "P-07"), ("SEASONAL", "P-08")):
            if val in flt.get("value", []):
                case(f"Excluded class ({val.lower()}) is not created", "Edge", pid,
                     f"{pid} has Employee_Class = {val}.",
                     f"1. Confirm {pid} appears in the source report.\n2. {run}\n3. Search {target_name} for {pid}'s Employee_ID.",
                     "No user is created. The record is skipped by the population filter.",
                     [] if val in built_vals else find_check(checks, category="Population filter"))
    case("Duplicate email across two workers", "Edge", "P-09 / P-10",
         "P-09 and P-10 share chris.nguyen@… with different Employee_IDs.",
         f"1. {run}\n2. Check the error log and both Employee_IDs in {target_name}.",
         "Records are NOT merged. Second record is rejected and logged for HRIS cleanup; first user unaffected.")
    case("Special characters are preserved", "Edge", "P-11",
         "P-11 has diacritics and a hyphenated surname.",
         f"1. {run}\n2. Open P-11 in {target_name} and in a report export.",
         "Name displays exactly as in source (Zoë Ó Briain-Søndergaard) in UI and CSV export.")
    case("Record missing a required field is rejected cleanly", "Edge", "P-12",
         "P-12 has no Work_Email.",
         f"1. {run}\n2. Check the integration error log.",
         "Record is skipped and logged with a readable reason. The rest of the batch completes.")
    if has("manager_email"):
        case("Manager change updates reporting line", "Edge", "P-13",
             "P-13 synced with manager lena.kovacs@….",
             f"1. Change P-13's manager in {src_name} to tomas.berg@….\n2. {run}\n3. Check P-13's manager in {target_name}.",
             "manager_email updates; P-13 moves to the new manager's team view.",
             find_check(checks, field="manager_email"))
    on_missing = parsed["target"].get("properties", {}).get("on_missing_from_feed") or (spec or {}).get("lifecycle", {}).get("missing_from_feed")
    if on_missing:
        case("Worker who drops off the feed", "Edge", "P-03",
             "A terminated worker ages out of the report window (include_terminated_days).",
             f"1. Remove P-03 from the source report (or wait past the window).\n2. {run}",
             f"User is handled per spec: {on_missing}. No error is raised.",
             find_check(checks, field="status"))
    case("Full-population reconciliation", "Core", "—",
         "First full run in the test tenant.",
         f"1. Count workers in the source report.\n2. Subtract excluded classes.\n3. Compare to active users in {target_name}.",
         "Counts reconcile exactly. Any difference is explained in the Issue Log.")
    return cases


def lms_cases(parsed, spec, checks):
    target_name = (spec or {}).get("target_system", "the LMS")
    cases = []

    def case(title, cat, persona, pre, steps, expected, checks_=None):
        cases.append({"title": title, "category": cat, "persona": persona, "pre": pre,
                      "steps": steps, "expected": expected, "checks": checks_ or []})

    run = "Run the completions sync (or wait for the next scheduled run)."
    case("Completion posts to learning history", "Core", "P-01",
         "P-01 exists as a student in the LMS; SAFE-101 exists in the LMS catalog as OLT.",
         f"1. P-01 completes SAFE-101 in the learning platform.\n2. {run}\n3. Open P-01's learning history in {target_name}.",
         "One completion record for SAFE-101 with correct date, status and score.")
    case("All mapped fields land correctly", "Core", "P-01",
         "Completion from the previous case has posted.",
         "Compare each field on the Field Mappings tab against the LMS record.",
         "Every field matches; mark each row Verified on the Field Mappings tab.",
         find_check(checks, category="Field mapping") + find_check(checks, category="Dates & time zones"))
    case("Late-evening completion keeps its calendar date", "Edge", "P-02",
         "P-02 completes at 11:41 PM local time.",
         f"1. P-02 completes COMP-220 at 23:41 America/Chicago.\n2. {run}\n3. Check COMPLETION_DATE in {target_name}.",
         "Completion date shows 2026-10-08, not 2026-10-09.",
         find_check(checks, category="Dates & time zones"))
    if parsed["filters"] or (spec or {}).get("filters"):
        case("Non-OLT course is not sent", "Edge", "P-03",
             "ILT-040 is an instructor-led course.",
             f"1. Mark P-03 complete on ILT-040.\n2. {run}",
             "No learning-history record is posted for ILT-040. Not logged as an error.",
             find_check(checks, category="Population filter"))
    case("Learner with no matching LMS student", "Edge", "P-04",
         "HF-20417 is a new hire not yet provisioned in the LMS.",
         f"1. P-04 completes SAFE-101.\n2. {run}\n3. Check the integration error log.",
         "Record is logged as 'student not found' and retried on the next run; no phantom student is created.")
    case("Duplicate completion is idempotent", "Edge", "P-05",
         "P-05's completion already posted.",
         f"1. Re-run the sync for the same window.\n2. Check P-05's history in {target_name}.",
         "Still exactly one record for that completion date.")
    case("Retake on a later date", "Edge", "P-06",
         "P-06 is the same learner as P-05, retaking SAFE-101 the next day.",
         f"1. Learner completes SAFE-101 again on 2026-10-09.\n2. {run}",
         "A second history record posts with the new date and score; the first is unchanged.")
    case("Course not in LMS catalog", "Edge", "P-07",
         "NEW-999 has no matching ITEM_ID in the LMS.",
         f"1. P-07 completes NEW-999.\n2. {run}\n3. Check the error log.",
         "Record is rejected with a readable 'item not found' reason; batch continues.")
    vol = (spec or {}).get("expected_volume", {}).get("peak_completions_per_run")
    case("Peak-volume batch completes", "Core", "—",
         f"Bulk-complete {vol or 'a large number of'} records in the test tenant (e.g. a compliance deadline).",
         f"1. Trigger a bulk completion batch.\n2. {run}\n3. Reconcile counts: platform completions vs LMS records.",
         "All records post. Rate-limit (HTTP 429) responses are retried, not dropped.",
         find_check(checks, category="Error handling"))
    return cases


# --------------------------------------------------------------------------------------
# 5. Workbook — "bold modern" styling
# --------------------------------------------------------------------------------------
FONT = "Arial"
INK, INK2, ACCENT, ACCENT_LIGHT = "0F172A", "1E293B", "4F46E5", "EEF2FF"
SURFACE, LINE, MUTED, ON_INK = "F8FAFC", "E2E8F0", "64748B", "CBD5E1"
INPUT = "FFF8DB"
# Status colors always ship with an icon + label, never color alone
GOOD, GOOD_BG = "0B7A0B", "DCF5DC"
CRIT, CRIT_BG = "B42318", "FDE3E1"
WARN_, WARN_BG = "8A5300", "FFF1C2"
INFO_, INFO_BG = "1D4ED8", "E0E9FF"

R = {"PASS": "✓ PASS", "FAIL": "✗ FAIL", "WARN": "⚠ WARN", "INFO": "ℹ INFO"}
T_NOT, T_PASS, T_FAIL, T_BLK = "○ Not run", "✓ Pass", "✗ Fail", "⚠ Blocked"
MET, NOT_MET = "✓ Met", "✗ Not met"
V_OK, V_BAD = "✓ Verified", "✗ Mismatch"
STATUS_STYLE = {  # label -> (text color, fill)
    R["PASS"]: (GOOD, GOOD_BG), R["FAIL"]: (CRIT, CRIT_BG), R["WARN"]: (WARN_, WARN_BG), R["INFO"]: (INFO_, INFO_BG),
    T_PASS: (GOOD, GOOD_BG), T_FAIL: (CRIT, CRIT_BG), T_BLK: (WARN_, WARN_BG), T_NOT: (MUTED, "F1F5F9"),
    MET: (GOOD, GOOD_BG), NOT_MET: (CRIT, CRIT_BG), V_OK: (GOOD, GOOD_BG), V_BAD: (CRIT, CRIT_BG),
}
TABLE_ROW = 6  # header row of every table sheet; data starts on row 7

HAIR = Side(style="thin", color=LINE)


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def f(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def col(n):
    return get_column_letter(n)


def band(ws, ncols, title, subtitle, legend=None):
    """Dark title band + accent strip across the sheet."""
    for c in range(1, ncols + 1):
        for r in (1, 2):
            ws.cell(row=r, column=c).fill = fill(INK)
        ws.cell(row=3, column=c).fill = fill(ACCENT)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 4
    ws["A1"] = title
    ws["A1"].font = f(18, True, "FFFFFF")
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws["A2"] = subtitle
    ws["A2"].font = f(10, False, ON_INK)
    ws["A2"].alignment = Alignment(vertical="top", indent=1)
    if legend:
        ws["A4"] = legend
        ws["A4"].font = f(9, False, MUTED, italic=True)
        ws["A4"].alignment = Alignment(indent=1)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ACCENT


def fit_rows(ws, first, last, widths, min_h=20):
    """Estimate wrapped row heights so nothing is clipped (Excel won't auto-fit rows written by openpyxl)."""
    for r in range(first, last + 1):
        lines = 1
        for c, w in enumerate(widths, 1):
            v = ws.cell(row=r, column=c).value
            if v is None or str(v).startswith("="):
                continue
            per_line = max(1, int(w * 1.3))
            n = sum(max(1, -(-len(par) // per_line)) for par in str(v).split("\n"))
            lines = max(lines, n)
        ws.row_dimensions[r].height = max(min_h, 13.5 * lines + 7)


def table(ws, headers, rows, widths, input_cols=(), center_cols=(), row=TABLE_ROW):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = f(10, True, "FFFFFF")
        c.fill = fill(INK2)
        c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    ws.row_dimensions[row].height = 26
    for r_i, r in enumerate(rows, row + 1):
        band_fill = fill(SURFACE) if (r_i - row) % 2 == 0 else None
        for c_i, v in enumerate(r, 1):
            c = ws.cell(row=r_i, column=c_i, value=v)
            c.font = f()
            c.alignment = Alignment(wrap_text=True, vertical="top", indent=0 if headers[c_i - 1] in center_cols else 1,
                                    horizontal="center" if headers[c_i - 1] in center_cols else None)
            c.border = Border(bottom=HAIR)
            if headers[c_i - 1] in input_cols:
                c.fill = fill(INPUT)
            elif band_fill:
                c.fill = band_fill
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[col(i)].width = w
    first, last = row + 1, row + len(rows)
    fit_rows(ws, first, last, widths)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return first, last


def dropdown(ws, rng, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)


def status_colors(ws, rng):
    for label, (fg, bg) in STATUS_STYLE.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{label}"'],
                                                      fill=fill(bg), font=Font(name=FONT, bold=True, color=fg)))


def section(ws, row, text, c1=2, c2=9):
    ws.cell(row=row, column=c1, value=text.upper()).font = f(10, True, ACCENT)
    for c in range(c1, c2 + 1):
        ws.cell(row=row, column=c).border = Border(bottom=Side(style="medium", color=ACCENT))
    ws.row_dimensions[row].height = 22


def build_workbook(parsed, spec, checks, cases, personas, out_path):
    wb = Workbook()
    customer = parsed["customer"]
    itype_label = {"hris_user_sync": "HRIS User Sync", "lms_completions": "LMS Completions"}.get(parsed["type"], parsed["type"])
    src = (spec or {}).get("source_system") or parsed["source"].get("connector", "Source")
    tgt = (spec or {}).get("target_system") or parsed["target"].get("connector", "Target")
    today = dt.date.today().isoformat()
    by_t = mapping_by_target(parsed)
    spec_req = {r["target"]: r for r in (spec or {}).get("required_mappings", [])}

    # Ranges other sheets point at (fixed layout: header on row 6, data from row 7)
    n_val = max(1, len(checks))
    VAL = f"'Pre-UAT Validation'!$E$7:$E${6 + n_val}"
    TC_S = f"'Test Cases'!$I$7:$I${6 + len(cases)}"
    TC_T = f"'Test Cases'!$C$7:$C${6 + len(cases)}"
    ISS_SEV, ISS_ST = "'Issue Log'!$C$8:$C$200", "'Issue Log'!$G$8:$G$200"
    SIGN = "'Sign-off'!$B$7:$B$11"

    # ---------- Pre-UAT Validation ----------
    ws_v = wb.active
    ws_v.title = "Pre-UAT Validation"
    band(ws_v, 8, "Pre-UAT Validation",
         f"{customer} · {itype_label} · the workflow as built vs. the signed requirements spec",
         "Resolve every ✗ FAIL before the customer starts testing. When a fix ships, re-run the generator (check IDs stay the same) or set Result and note it in Resolution.")
    if checks:
        v_rows = [[c["id"], c["category"], c["requirement"], c["as_built"], R[c["result"]],
                   c["severity"], c["action"], ""] for c in checks]
    else:
        v_rows = [["—", "—", "No requirements spec supplied", "—", R["INFO"], "—",
                   "Provide the signed spec to validate the build before UAT.", ""]]
    v_first, v_last = table(ws_v, ["Check", "Category", "Signed requirement", "Workflow as built", "Result",
                                   "Severity", "Recommended action", "Resolution"],
                            v_rows, [9, 18, 42, 36, 12, 10, 52, 30], input_cols=("Resolution",),
                            center_cols=("Check", "Result", "Severity"))
    dropdown(ws_v, f"E{v_first}:E{v_last}", list(R.values()))
    status_colors(ws_v, f"E{v_first}:E{v_last}")
    ws_v.conditional_formatting.add(f"F{v_first}:F{v_last}", CellIsRule(operator="equal", formula=['"High"'],
                                    font=Font(name=FONT, bold=True, color=CRIT)))
    ws_v.sheet_properties.tabColor = CRIT if any(c["result"] == "FAIL" for c in checks) else GOOD

    # ---------- Field Mappings ----------
    ws_m = wb.create_sheet("Field Mappings")
    band(ws_m, 9, "Field Mappings", f"Every field written to {tgt}, plus agreed fields not yet built",
         "Customer: check each field against one real record and set Verified. Sample values come from persona P-02 on Test Data.")
    sample = personas[1][1] if len(personas[1]) > 1 else {}
    open_by_field = {}
    for c in checks:
        if c["field"] and c["result"] != "PASS":
            open_by_field.setdefault(c["field"], []).append(c)
    m_rows, m_fill = [], []
    all_targets = list(by_t) + [t for t in spec_req if t not in by_t]
    for i, t in enumerate(all_targets, 1):
        m = by_t.get(t)
        built_src = m["source"] if m else "— not built —"
        agreed_src = spec_req[t]["source"] if t in spec_req else "— not in spec —"
        transform = (m or {}).get("transform") or "—"
        if (m or {}).get("value_map"):
            transform += ": " + ", ".join(f"{k}→{v}" for k, v in m["value_map"].items())
        if (m or {}).get("timezone"):
            transform += f" (tz {m['timezone']})"
        flags = open_by_field.get(t, [])
        worst = "FAIL" if any(c["result"] == "FAIL" for c in flags) else "WARN" if any(
            c["result"] == "WARN" for c in flags) else "INFO" if flags else None
        flag = (R[worst][0] + " " + ", ".join(c["id"] for c in flags)) if worst else ""
        sample_val = sample.get(m["source"] if m else spec_req[t]["source"], "")
        m_rows.append([i, built_src, t, agreed_src, transform, sample_val, flag, "", ""])
        m_fill.append(worst)
    m_first, m_last = table(ws_m, ["#", f"Source field ({src})", f"Target field ({tgt})", "Agreed source (spec)",
                                   "Transform", "Sample value", "Flag", "Verified", "Notes"],
                            m_rows, [5, 26, 22, 24, 34, 28, 14, 14, 28], input_cols=("Verified", "Notes"),
                            center_cols=("#", "Flag"))
    for r, worst in zip(range(m_first, m_last + 1), m_fill):
        if worst:
            fg, bg = STATUS_STYLE[R[worst]]
            cell = ws_m.cell(row=r, column=7)
            cell.fill, cell.font = fill(bg), f(10, True, fg)
    dropdown(ws_m, f"H{m_first}:H{m_last}", [V_OK, V_BAD, "N/A"])
    status_colors(ws_m, f"H{m_first}:H{m_last}")

    # ---------- Test Cases ----------
    ws_tc = wb.create_sheet("Test Cases")
    band(ws_tc, 13, "Test Cases", f"{customer} · {itype_label} · {src} → {tgt}",
         "Yellow cells are for the tester. Set Status from the dropdown; log every ✗ Fail or ⚠ Blocked in the Issue Log and put its ID here.")
    tc_rows = []
    for i, c in enumerate(cases, 1):
        dep = ""
        if c["checks"]:
            ids = ", ".join(x["id"] for x in c["checks"])
            dep = f"⚠ Will fail until {ids} {'is' if len(c['checks']) == 1 else 'are'} fixed"
        tc_rows.append([f"TC-{i:02d}", c["title"], c["category"], c["persona"], c["pre"], c["steps"],
                        c["expected"], "", T_NOT, "", "", "", dep])
    tc_first, tc_last = table(ws_tc, ["TC ID", "Scenario", "Type", "Test data", "Preconditions", "Steps",
                                      "Expected result", "Actual result", "Status", "Tester", "Date", "Issue ID",
                                      "Pre-UAT dependency"],
                              tc_rows, [8, 28, 8, 10, 32, 44, 38, 28, 13, 12, 11, 10, 28],
                              input_cols=("Actual result", "Status", "Tester", "Date", "Issue ID"),
                              center_cols=("TC ID", "Type", "Test data", "Status"))
    dropdown(ws_tc, f"I{tc_first}:I{tc_last}", [T_NOT, T_PASS, T_FAIL, T_BLK])
    status_colors(ws_tc, f"I{tc_first}:I{tc_last}")
    ws_tc.conditional_formatting.add(f"C{tc_first}:C{tc_last}", CellIsRule(
        operator="equal", formula=['"Core"'], fill=fill(ACCENT_LIGHT), font=Font(name=FONT, bold=True, color=ACCENT)))
    for r in range(tc_first, tc_last + 1):
        if ws_tc.cell(row=r, column=13).value:
            ws_tc.cell(row=r, column=13).font = f(9, True, WARN_)

    # ---------- Test Data ----------
    ws_d = wb.create_sheet("Test Data")
    cols, rows = personas
    band(ws_d, len(cols), "Test Data", "Fictional personas the test cases use",
         "Create or modify these records in the TEST tenant only. Times are the customer's local time zone.")
    table(ws_d, cols, [[p[c] for c in cols] for p in rows], [8, 24] + [19] * (len(cols) - 2),
          center_cols=("Persona",))

    # ---------- Issue Log ----------
    ws_i = wb.create_sheet("Issue Log")
    band(ws_i, 9, "Issue Log", "Every failed or blocked test case gets a row here",
         "Row 7 is a greyed-out example showing the format; it is not counted. Log issues from row 8.")
    i_headers = ["Issue ID", "TC ID", "Severity", "Description", "Owner", "Opened", "Status", "Resolution", "Closed"]
    example = ["ISS-000", "TC-06", "High", "Rehired worker created as a new user instead of reactivating the original.",
               "Integration team", today, "Example", "Hire date source corrected; re-tested OK.", today]
    blank = [[f"ISS-{n:03d}", "", "", "", "", "", "", "", ""] for n in range(1, 31)]
    i_first, i_last = table(ws_i, i_headers, [example] + blank, [10, 8, 10, 50, 18, 11, 13, 40, 11],
                            input_cols=tuple(i_headers[1:]), center_cols=("Issue ID", "TC ID", "Severity", "Status"))
    for c in range(1, 10):
        cell = ws_i.cell(row=i_first, column=c)
        cell.font = f(9, italic=True, color="94A3B8")
        cell.fill = fill("F1F5F9")
    dropdown(ws_i, f"C{i_first + 1}:C{i_last}", ["High", "Medium", "Low"])
    dropdown(ws_i, f"G{i_first + 1}:G{i_last}", ["Open", "In progress", "Resolved", "Won't fix"])
    ws_i.conditional_formatting.add(f"C{i_first + 1}:C{i_last}", CellIsRule(
        operator="equal", formula=['"High"'], font=Font(name=FONT, bold=True, color=CRIT)))
    ws_i.conditional_formatting.add(f"G{i_first + 1}:G{i_last}", CellIsRule(
        operator="equal", formula=['"Open"'], fill=fill(CRIT_BG), font=Font(name=FONT, bold=True, color=CRIT)))
    ws_i.conditional_formatting.add(f"G{i_first + 1}:G{i_last}", CellIsRule(
        operator="equal", formula=['"Resolved"'], fill=fill(GOOD_BG), font=Font(name=FONT, bold=True, color=GOOD)))

    # ---------- Sign-off ----------
    ws_s = wb.create_sheet("Sign-off")
    band(ws_s, 4, "UAT Sign-off", f"{customer} · {itype_label}",
         "Criteria are formulas and update as the workbook is filled in. Sign only when every criterion reads ✓ Met.")
    n_map = m_last - m_first + 1
    MAP = f"'Field Mappings'!$H${m_first}:$H${m_last}"
    crit = [
        ("No failing pre-UAT checks", f'=IF(COUNTIF({VAL},"{R["FAIL"]}")=0,"{MET}","{NOT_MET}")'),
        ("Every Core test case passed", f'=IF(COUNTIFS({TC_T},"Core",{TC_S},"<>{T_PASS}")=0,"{MET}","{NOT_MET}")'),
        ("Every Edge test case run, nothing blocked", f'=IF(COUNTIFS({TC_T},"Edge",{TC_S},"{T_NOT}")+COUNTIFS({TC_T},"Edge",{TC_S},"{T_BLK}")=0,"{MET}","{NOT_MET}")'),
        ("No open High-severity issues", f'=IF(COUNTIFS({ISS_SEV},"High",{ISS_ST},"Open")=0,"{MET}","{NOT_MET}")'),
        ("Every field mapping verified", f'=IF(COUNTIF({MAP},"{V_OK}")+COUNTIF({MAP},"N/A")={n_map},"{MET}","{NOT_MET}")'),
    ]
    s_first, s_last = table(ws_s, ["Criterion", "Status"], [[k, v] for k, v in crit], [48, 16, 26, 14],
                            center_cols=("Status",))
    status_colors(ws_s, f"B{s_first}:B{s_last}")
    r = s_last + 2
    section(ws_s, r, "Approvals", 1, 4)
    signers = [[s["name"], s["role"], "", ""] for s in (spec or {}).get("signed_off", [])]
    signers.append(["", "Integration Engineer", "", ""])
    table(ws_s, ["Name", "Role", "Signature", "Date"], signers, [48, 40, 26, 14],
          input_cols=("Signature", "Date"), row=r + 1)
    ws_s.freeze_panes = None

    # ---------- AI Instructions ----------
    ws_a = wb.create_sheet("AI Instructions")
    band(ws_a, 2, "AI Instructions", "For Claude (or any agent) maintaining this workbook. Humans can skip this tab.")
    instr = [
        ("Source of truth", "The workflow export and the signed requirements spec. Never invent fields, values or customer decisions."),
        ("Regenerating", "When the workflow changes, re-run generate_uat.py with the new export. Check IDs are stable between runs. Carry over Actual result, Status, Tester, Date and Issue ID by matching on Scenario, not TC ID."),
        ("Validation", "A ✗ FAIL on Pre-UAT Validation blocks customer testing. Summarise each FAIL in plain language with the fix, ordered by severity. ⚠ WARN and ℹ INFO items need a customer decision, not a silent fix."),
        ("Test cases", "Add a case only when the workflow or spec introduces a behaviour not already covered. Each case needs a persona from Test Data, numbered steps and one observable expected result."),
        ("Test data", "All personas are fictional and live in the TEST tenant. Never copy production records into this workbook."),
        ("Tone", "Customer-facing. Short sentences, no internal jargon, no blame. Name systems the way the customer does."),
        ("Sign-off", "Do not type ✓ Met by hand. Criteria are formulas; fix the underlying data instead."),
    ]
    table(ws_a, ["Topic", "Instruction"], [list(x) for x in instr], [22, 110])

    # ---------- Overview (dashboard) ----------
    ws = wb.create_sheet("Overview", 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = INK
    ws.column_dimensions["A"].width = 3
    for c in range(2, 10):
        ws.column_dimensions[col(c)].width = 15
    ws.column_dimensions["J"].width = 3
    for r in (1, 2, 3):
        for c in range(1, 11):
            ws.cell(row=r, column=c).fill = fill(INK)
    for c in range(1, 11):
        ws.cell(row=4, column=c).fill = fill(ACCENT)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 40
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 5
    ws["B1"] = f"UAT WORKBOOK  ·  {itype_label.upper()}"
    ws["B1"].font = f(9, True, "A5B4FC")
    ws["B1"].alignment = Alignment(vertical="bottom")
    ws["B2"] = customer
    ws["B2"].font = f(26, True, "FFFFFF")
    ws["B2"].alignment = Alignment(vertical="center")
    ws["B3"] = f"{src}  →  {tgt}   ·   {parsed['env'].upper()}   ·   generated {today}"
    ws["B3"].font = f(10, False, ON_INK)
    ws["B3"].alignment = Alignment(vertical="top")

    fails = f'COUNTIF({VAL},"{R["FAIL"]}")'
    total, notrun = f"COUNTA({TC_S})", f'COUNTIF({TC_S},"{T_NOT}")'
    passed, failed, blocked = (f'COUNTIF({TC_S},"{T_PASS}")', f'COUNTIF({TC_S},"{T_FAIL}")',
                               f'COUNTIF({TC_S},"{T_BLK}")')

    # Readiness banner — three states, driven entirely by formulas
    ws.merge_cells("B6:I7")
    ws["B6"] = (f'=IF({fails}>0,"✗   NOT READY FOR CUSTOMER UAT   ·   "&{fails}&" blocking check"&IF({fails}=1,"","s"),'
                f'IF(COUNTIF({SIGN},"{MET}")=5,"✓   READY FOR SIGN-OFF   ·   every criterion met",'
                f'"▶   READY FOR CUSTOMER UAT   ·   "&{notrun}&" of "&{total}&" test cases to run"))')
    ws["B6"].font = f(15, True, "FFFFFF")
    ws["B6"].alignment = Alignment(horizontal="center", vertical="center")
    ws["B6"].fill = fill(ACCENT)
    ws.row_dimensions[6].height = 26
    ws.row_dimensions[7].height = 26
    ws.conditional_formatting.add("B6:I7", FormulaRule(formula=[f"{fails}>0"], fill=fill("D03B3B"),
                                                       font=Font(name=FONT, size=15, bold=True, color="FFFFFF")))
    ws.conditional_formatting.add("B6:I7", FormulaRule(formula=[f'COUNTIF({SIGN},"{MET}")=5'], fill=fill(GOOD),
                                                       font=Font(name=FONT, size=15, bold=True, color="FFFFFF")))

    # KPI tiles
    tiles = [
        ("Blocking checks", f"={fails}", '="must be 0 before UAT starts"', INK2),
        ("Test cases", f"={total}", f'=COUNTIF({TC_T},"Core")&" core  ·  "&COUNTIF({TC_T},"Edge")&" edge"', ACCENT),
        ("Passed", f"={passed}", f'={failed}&" failed  ·  "&{blocked}&" blocked"', ACCENT),
        ("Open issues", f'=COUNTIF({ISS_ST},"Open")', '="on the Issue Log"', ACCENT),
    ]
    ws.row_dimensions[9].height = 5
    ws.row_dimensions[10].height = 22
    ws.row_dimensions[11].height = 46
    ws.row_dimensions[12].height = 22
    for i, (label, value, caption, top) in enumerate(tiles):
        c1, c2 = 2 + i * 2, 3 + i * 2
        for c in (c1, c2):
            ws.cell(row=9, column=c).fill = fill(top)
            for r in (10, 11, 12):
                ws.cell(row=r, column=c).fill = fill(SURFACE)
        for r, v, font in ((10, label.upper(), f(9, True, MUTED)), (11, value, f(30, True, INK)), (12, caption, f(9, False, MUTED))):
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
            cell = ws.cell(row=r, column=c1, value=v)
            cell.font = font
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.conditional_formatting.add("B11:C11", CellIsRule(operator="greaterThan", formula=["0"],
                                                        font=Font(name=FONT, size=30, bold=True, color=CRIT)))
    ws.conditional_formatting.add("B11:C11", CellIsRule(operator="equal", formula=["0"],
                                                        font=Font(name=FONT, size=30, bold=True, color=GOOD)))

    # Test progress bar (pure formula, renders in Excel, Sheets and LibreOffice)
    section(ws, 14, "Test progress")
    ws.merge_cells("B15:H15")
    done = f"({total}-{notrun})"
    ws["B15"] = f'=IF({total}=0,"",REPT("█",ROUND({done}/{total}*60,0))&REPT("░",60-ROUND({done}/{total}*60,0)))'
    ws["B15"].font = f(12, False, ACCENT)
    ws["B15"].alignment = Alignment(vertical="center")
    ws["I15"] = f'=IF({total}=0,"",TEXT({done}/{total},"0%")&" run")'
    ws["I15"].font = f(13, True, INK)
    ws["I15"].alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[15].height = 26
    ws.merge_cells("B16:I16")
    ws["B16"] = f'="✓ "&{passed}&" passed     ✗ "&{failed}&" failed     ⚠ "&{blocked}&" blocked     ○ "&{notrun}&" not run"'
    ws["B16"].font = f(9, False, MUTED)

    # Details
    section(ws, 18, "Details")
    ws.row_dimensions[19].height = 6
    signed = ("\n".join(f"{s['name']}, {s['role']} ({s['date']})" for s in spec.get("signed_off", []))
              if spec else "Not supplied — validation skipped")
    meta = [
        ("Workflow", parsed["title"]), ("Workflow ID", parsed["id"]),
        ("Schedule (as built)", human_schedule(parsed["schedule"])),
        ("Match key", parsed["target"].get("properties", {}).get("match_on")
         or " + ".join(parsed["target"].get("properties", {}).get("dedupe_on", [])) or "—"),
        ("Spec signed by", signed),
    ]
    r = 20
    for k, v in meta:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=9)
        ws.cell(row=r, column=2, value=k).font = f(10, True, MUTED)
        ws.cell(row=r, column=2).alignment = Alignment(vertical="top")
        ws.cell(row=r, column=4, value=v).font = f(10)
        ws.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 15 * (str(v).count("\n") + 1) + 5
        r += 1

    # How to use
    r += 1
    section(ws, r, "How to use this workbook")
    r += 1
    ws.row_dimensions[r].height = 6
    steps = [
        ("1", "Pre-UAT Validation", "Integration team fixes every ✗ FAIL before the customer starts testing."),
        ("2", "Field Mappings", "Customer checks each field against a real record and marks it ✓ Verified."),
        ("3", "Test Cases", "Run each case with its persona from Test Data; record the result and status."),
        ("4", "Issue Log", "Log every failed or blocked case and link the Issue ID back to the test case."),
        ("5", "Sign-off", "When every criterion reads ✓ Met, both sides sign and cutover is scheduled."),
    ]
    for n, tab, text in steps:
        r += 1
        ws.cell(row=r, column=2, value=f"{n}  {tab}").font = f(10, True, ACCENT)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=9)
        ws.cell(row=r, column=4, value=text).font = f(10)
        ws.row_dimensions[r].height = 20
    r += 2
    ws.cell(row=r, column=2, value="Yellow cells are inputs. Everything else is generated or calculated.").font = f(9, False, MUTED, italic=True)
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    order = ["Overview", "Pre-UAT Validation", "Field Mappings", "Test Cases", "Test Data",
             "Issue Log", "Sign-off", "AI Instructions"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    for w in wb.worksheets:
        w.sheet_view.tabSelected = w.title == "Overview"
    wb.save(out_path)


# --------------------------------------------------------------------------------------
# 6. CLI
# --------------------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workflow", required=True, help="Workflow export JSON (Tray-style)")
    ap.add_argument("--spec", help="Signed requirements spec JSON (enables pre-UAT validation)")
    ap.add_argument("--out", help="Output .xlsx path")
    args = ap.parse_args()

    def load(path, what):
        try:
            return json.loads(Path(path).read_text(encoding="utf-8"))
        except FileNotFoundError:
            sys.exit(f"Error: {what} file not found: {path}")
        except json.JSONDecodeError as e:
            sys.exit(f"Error: {what} is not valid JSON ({path}, line {e.lineno}): {e.msg}")

    raw = load(args.workflow, "workflow")
    if not isinstance(raw, dict) or not (raw.get("workflow", raw)).get("steps"):
        sys.exit("Error: workflow export has no steps. Expected {\"workflow\": {\"steps\": [...]}} — see the input format reference.")
    parsed = parse_workflow(raw)
    spec = load(args.spec, "requirements spec") if args.spec else None
    checks = validate(parsed, spec)

    if parsed["type"] == "hris_user_sync":
        cases, personas = hris_cases(parsed, spec, checks), hris_personas(parsed["customer"])
    elif parsed["type"] == "lms_completions":
        cases, personas = lms_cases(parsed, spec, checks), lms_personas(parsed["customer"])
    else:
        sys.exit(f"Error: unsupported integration type (source connector '{parsed['source'].get('connector')}', "
                 f"target '{parsed['target'].get('connector')}'). Supported: hris_user_sync, lms_completions — "
                 f"set one in the workflow's tags if the connectors aren't recognised.")

    slug = re.sub(r"[^A-Za-z0-9]+", "-", parsed["customer"]).strip("-")
    out = args.out or f"UAT_{slug}_{parsed['type']}_{dt.date.today().isoformat()}.xlsx"
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    build_workbook(parsed, spec, checks, cases, personas, out)

    summary = {
        "workbook": str(Path(out).resolve()),
        "customer": parsed["customer"],
        "integration_type": parsed["type"],
        "workflow": parsed["title"],
        "spec_supplied": spec is not None,
        "verdict": ("no spec — not validated" if spec is None else
                    "not ready" if any(c["result"] == "FAIL" for c in checks) else
                    "ready, with warnings to review" if any(c["result"] == "WARN" for c in checks) else
                    "ready for customer UAT"),
        "validation": {k: sum(1 for c in checks if c["result"] == k) for k in ("PASS", "FAIL", "WARN", "INFO")},
        "issues": [{k: c[k] for k in ("id", "result", "severity", "category", "requirement", "as_built", "action")}
                   for c in checks if c["result"] in ("FAIL", "WARN", "INFO")],
        "test_cases": {"total": len(cases), "core": sum(c["category"] == "Core" for c in cases),
                       "edge": sum(c["category"] == "Edge" for c in cases),
                       "expected_to_fail_until_fixed": sum(1 for c in cases if c["checks"])},
    }
    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
